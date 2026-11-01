"""Metadata policy, metadata change count, and coverage in HTML/CSV/Excel diff reports."""

from __future__ import annotations

import csv
import logging
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cja_auto_sdr.diff.models import DiffCoverage, DiffResult, DiffSummary, MetadataDiff, coverage_rows
from cja_auto_sdr.output.diff.csv import write_diff_csv_output
from cja_auto_sdr.output.diff.excel import write_diff_excel_output
from cja_auto_sdr.output.diff.html import write_diff_html_output

logger = logging.getLogger("test_diff_report_coverage")


def _coverage(policy: str = "reported") -> DiffCoverage:
    return DiffCoverage(
        compare_fields=["name", "description"],
        ignored_fields=[],
        component_types=["metrics", "dimensions"],
        inventories=[],
        metadata_policy=policy,
        show_only=None,
        limitations=["Only listed fields compared", "<b>x</b>"],
    )


def _result(*, metadata_changes: int, include: bool, coverage: DiffCoverage | None) -> DiffResult:
    return DiffResult(
        summary=DiffSummary(metadata_changes=metadata_changes, include_metadata=include),
        metadata_diff=MetadataDiff(source_name="A", target_name="B", source_id="dv_a", target_id="dv_b"),
        metric_diffs=[],
        dimension_diffs=[],
        coverage=coverage,
    )


def _csv_rows(path: str) -> dict[str, str]:
    with open(path, newline="", encoding="utf-8") as f:
        return {r["Property"]: r["Value"] for r in csv.DictReader(f)}


def _sheet_rows(path: str, name: str) -> dict[str, str]:
    ws = load_workbook(path)[name]
    return {str(r[0]): ("" if r[1] is None else str(r[1])) for r in ws.iter_rows(min_row=2, values_only=True)}


CASES = [
    (1, False, "reported"),
    (1, True, "included"),
    (0, False, "reported"),
    (0, True, "included"),
]


@pytest.mark.parametrize(("count", "include", "policy"), CASES)
def test_html_metadata_line_and_policy(tmp_path, count, include, policy):
    r = _result(metadata_changes=count, include=include, coverage=_coverage(policy))
    content = Path(write_diff_html_output(r, "d", str(tmp_path), logger)).read_text(encoding="utf-8")
    assert f"Metadata policy: {policy}" in content
    if count:
        suffix = "(counted)" if include else "(not counted; use --include-metadata-changes)"
        assert f"Metadata changes: 1 {suffix}" in content
    else:
        assert "Metadata changes:" not in content


def test_html_coverage_section(tmp_path):
    r = _result(metadata_changes=1, include=False, coverage=_coverage())
    content = Path(write_diff_html_output(r, "d", str(tmp_path), logger)).read_text(encoding="utf-8")
    assert "Comparison coverage" in content
    assert "compare_fields" in content
    assert "name, description" in content
    assert "limitations" in content
    assert "Only listed fields compared, &lt;b&gt;x&lt;/b&gt;" in content
    assert "<b>x</b>" not in content
    assert "—" in content  # show_only None / empty ignored_fields


def test_html_no_coverage(tmp_path):
    r = _result(metadata_changes=0, include=False, coverage=None)
    content = Path(write_diff_html_output(r, "d", str(tmp_path), logger)).read_text(encoding="utf-8")
    assert "Comparison coverage" not in content
    assert "Metadata policy: reported" in content


@pytest.mark.parametrize(("count", "include", "policy"), CASES)
def test_csv_metadata_rows(tmp_path, count, include, policy):
    r = _result(metadata_changes=count, include=include, coverage=_coverage(policy))
    d = write_diff_csv_output(r, "d", str(tmp_path), logger)
    rows = _csv_rows(os.path.join(d, "metadata.csv"))
    assert rows["metadata_changes"] == str(count)
    assert rows["metadata_policy"] == policy
    assert list(rows).index("metadata_changes") == list(rows).index("total_changes") + 1


def test_csv_coverage_file(tmp_path):
    r = _result(metadata_changes=1, include=False, coverage=_coverage())
    d = write_diff_csv_output(r, "d", str(tmp_path), logger)
    rows = _csv_rows(os.path.join(d, "coverage.csv"))
    assert rows["compare_fields"] == "name; description"
    assert rows["limitations"] == "Only listed fields compared; <b>x</b>"
    assert rows["metadata_policy"] == "reported"


def test_csv_no_coverage(tmp_path):
    r = _result(metadata_changes=0, include=False, coverage=None)
    d = write_diff_csv_output(r, "d", str(tmp_path), logger)
    assert not os.path.exists(os.path.join(d, "coverage.csv"))
    assert _csv_rows(os.path.join(d, "metadata.csv"))["metadata_policy"] == "reported"


@pytest.mark.parametrize(("count", "include", "policy"), CASES)
def test_excel_metadata_rows(tmp_path, count, include, policy):
    r = _result(metadata_changes=count, include=include, coverage=_coverage(policy))
    p = write_diff_excel_output(r, "d", str(tmp_path), logger)
    rows = _sheet_rows(p, "Metadata")
    assert rows["Metadata Changes"] == str(count)
    assert rows["Metadata Policy"] == policy


def test_excel_coverage_sheet(tmp_path):
    r = _result(metadata_changes=1, include=False, coverage=_coverage())
    p = write_diff_excel_output(r, "d", str(tmp_path), logger)
    wb = load_workbook(p)
    assert wb.sheetnames.index("Coverage") == wb.sheetnames.index("Metadata") + 1
    rows = _sheet_rows(p, "Coverage")
    assert rows["compare_fields"] == "name; description"
    assert "Only listed fields compared" in rows["limitations"]


def test_excel_no_coverage(tmp_path):
    r = _result(metadata_changes=0, include=False, coverage=None)
    p = write_diff_excel_output(r, "d", str(tmp_path), logger)
    wb = load_workbook(p)
    assert "Coverage" not in wb.sheetnames
    assert _sheet_rows(p, "Metadata")["Metadata Policy"] == "reported"


def test_summary_metadata_policy_property():
    assert DiffSummary().metadata_policy == "reported"
    assert DiffSummary(include_metadata=True).metadata_policy == "included"


def test_coverage_rows_helper():
    rows = dict(coverage_rows(_coverage()))
    assert rows["compare_fields"] == "name; description"
    assert rows["ignored_fields"] == ""
    assert rows["show_only"] == ""
    assert rows["metadata_policy"] == "reported"


_OLD_DESC = "Counts visits from the EMEA region only"
_NEW_DESC = "Counts visits from the EMEA region and APAC"


def _long_change_result() -> DiffResult:
    from cja_auto_sdr.diff.models import ChangeType, ComponentDiff

    diff = ComponentDiff(
        id="metrics/a",
        name="A",
        change_type=ChangeType.MODIFIED,
        changed_fields={"description": (_OLD_DESC, _NEW_DESC)},
    )
    return DiffResult(
        summary=DiffSummary(),
        metadata_diff=MetadataDiff(source_name="A", target_name="B", source_id="dv_a", target_id="dv_b"),
        metric_diffs=[diff],
        dimension_diffs=[],
    )


def test_file_outputs_keep_full_change_values(tmp_path):
    from cja_auto_sdr.output.diff.markdown import write_diff_markdown_output

    r = _long_change_result()
    csv_dir = write_diff_csv_output(r, "d", str(tmp_path / "csv"), logger)
    csv_text = "".join(Path(csv_dir, p).read_text(encoding="utf-8") for p in os.listdir(csv_dir) if p.endswith(".csv"))
    assert _OLD_DESC in csv_text
    assert _NEW_DESC in csv_text

    xlsx = write_diff_excel_output(r, "d", str(tmp_path / "xlsx"), logger)
    wb = load_workbook(xlsx)
    cells = [str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c is not None]
    assert any(_OLD_DESC in c and _NEW_DESC in c for c in cells)

    html = Path(write_diff_html_output(r, "d", str(tmp_path / "html"), logger)).read_text(encoding="utf-8")
    assert _OLD_DESC in html
    assert _NEW_DESC in html

    md = Path(write_diff_markdown_output(r, "d", str(tmp_path / "md"), logger)).read_text(encoding="utf-8")
    assert _OLD_DESC in md
    assert _NEW_DESC in md


_INV_OLD = "Visits from the EMEA region only"
_INV_NEW = "Visits from the EMEA region and APAC"


@pytest.mark.parametrize("fmt", ["csv", "excel", "html", "markdown"])
def test_inventory_file_outputs_keep_full_change_values(tmp_path, fmt):
    from cja_auto_sdr.diff.models import ChangeType, InventoryItemDiff
    from cja_auto_sdr.output.diff.markdown import write_diff_markdown_output

    item = InventoryItemDiff(
        id="s_mod",
        name="Mod seg",
        change_type=ChangeType.MODIFIED,
        inventory_type="segment",
        changed_fields={"description": (_INV_OLD, _INV_NEW)},
    )
    r = DiffResult(
        summary=DiffSummary(segments_modified=1),
        metadata_diff=MetadataDiff(source_name="A", target_name="B", source_id="dv_a", target_id="dv_b"),
        metric_diffs=[],
        dimension_diffs=[],
        segments_diffs=[item],
    )
    out = str(tmp_path / fmt)
    if fmt == "csv":
        d = write_diff_csv_output(r, "d", out, logger)
        text = "".join(Path(d, p).read_text(encoding="utf-8") for p in os.listdir(d) if p.endswith(".csv"))
    elif fmt == "excel":
        wb = load_workbook(write_diff_excel_output(r, "d", out, logger))
        text = "\n".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c is not None)
    elif fmt == "html":
        text = Path(write_diff_html_output(r, "d", out, logger)).read_text(encoding="utf-8")
    else:
        text = Path(write_diff_markdown_output(r, "d", out, logger)).read_text(encoding="utf-8")
    assert _INV_OLD in text
    assert _INV_NEW in text


def test_console_truncation_marks_ellipsis():
    from cja_auto_sdr.output.diff.common import _format_diff_value

    assert _format_diff_value("x" * 40) == "x" * 30 + "..."
    assert _format_diff_value("x" * 30) == "x" * 30


def _inventory_result() -> DiffResult:
    from cja_auto_sdr.diff.models import ChangeType, InventoryItemDiff

    def item(i, n, ct, kind, fields=None):
        return InventoryItemDiff(id=i, name=n, change_type=ct, inventory_type=kind, changed_fields=fields)

    segs = [
        item("s_new", "New seg", ChangeType.ADDED, "segment"),
        item("s_old", "Old seg", ChangeType.REMOVED, "segment"),
        item("s_mod", "Mod seg", ChangeType.MODIFIED, "segment", {"description": ("a", "b")}),
    ]
    cms = [item("cm_new", "New cm", ChangeType.ADDED, "calculated_metric")]
    return DiffResult(
        summary=DiffSummary(
            segments_added=1,
            segments_removed=1,
            segments_modified=1,
            calc_metrics_added=1,
        ),
        metadata_diff=MetadataDiff(source_name="A", target_name="B", source_id="dv_a", target_id="dv_b"),
        metric_diffs=[],
        dimension_diffs=[],
        calc_metrics_diffs=cms,
        segments_diffs=segs,
    )


def test_grouped_lists_inventory_changes():
    from cja_auto_sdr.output.diff.grouped import write_diff_grouped_by_field_output

    out = write_diff_grouped_by_field_output(_inventory_result(), use_color=False)
    seg_section = out.split("Segments")[1]
    for text in ("s_new", "New seg", "s_old", "Old seg", "s_mod", "Mod seg"):
        assert text in seg_section
    assert "Calculated Metrics" in out
    assert "cm_new" in out.split("Calculated Metrics")[1]
    assert "Total changes: 4" in out


def test_grouped_inventory_follows_limit():
    from cja_auto_sdr.output.diff.grouped import write_diff_grouped_by_field_output

    out = write_diff_grouped_by_field_output(_inventory_result(), use_color=False, limit=1)
    assert "... and 2 more" in out.split("Segments")[1]
    unlimited = write_diff_grouped_by_field_output(_inventory_result(), use_color=False, limit=0)
    assert "more" not in unlimited.split("Segments")[1]
